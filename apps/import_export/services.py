import csv
from io import StringIO
from pathlib import Path

from django.http import HttpResponse
from django.utils.text import slugify

from apps.catalog.models import Author, Book, Category, Publisher, normalize_isbn
from apps.circulation.models import BorrowRecord
from apps.fines.models import Fine
from apps.members.models import Member
from apps.reports.services import export_basic_xlsx


BOOK_HEADERS = ["title", "isbn", "book_code", "author", "publisher", "category", "publication_year", "language"]


def unique_category_code(name):
    base_code = (slugify(name).replace("-", "_").upper()[:20] or "CATEGORY")
    code = base_code
    index = 1
    while Category.objects.filter(code=code).exists():
        index += 1
        suffix = f"_{index}"
        code = f"{base_code[: 30 - len(suffix)]}{suffix}"
    return code


def import_books_from_job(job):
    path = Path(job.source_file.path)
    try:
        rows = read_rows(path)
        success = 0
        failed = 0
        for raw_row in rows:
            row = {str(key).strip(): (value.strip() if isinstance(value, str) else value) for key, value in raw_row.items() if key}
            if not row.get("title"):
                failed += 1
                continue
            try:
                publisher = None
                if row.get("publisher"):
                    publisher = Publisher.objects.filter(name__iexact=row["publisher"]).first()
                    if not publisher:
                        publisher = Publisher.objects.create(name=row["publisher"])
                category = None
                if row.get("category"):
                    category, _ = Category.objects.get_or_create(
                        name=row["category"],
                        defaults={"code": unique_category_code(row["category"])},
                    )
                isbn = normalize_isbn(row.get("isbn"))
                defaults = {
                    "title": row["title"],
                    "book_code": row.get("book_code") or "",
                    "publisher": publisher,
                    "category": category,
                    "publication_year": row.get("publication_year") or None,
                    "language": row.get("language") or "English",
                }
                if isbn:
                    book, _ = Book.objects.get_or_create(isbn=isbn, defaults=defaults)
                else:
                    book = Book.objects.create(**defaults)
                author_name = row.get("author") or row.get("authors")
                if author_name:
                    author = Author.objects.filter(name__iexact=author_name).first()
                    if not author:
                        author = Author.objects.create(name=author_name)
                    book.authors.add(author)
                success += 1
            except Exception:
                failed += 1
                continue
        job.total_rows = len(rows)
        job.success_rows = success
        job.failed_rows = failed
        job.status = job.COMPLETED
        job.message = "Import completed."
    except Exception as exc:
        job.status = job.FAILED
        job.message = str(exc)
    job.save(update_fields=["total_rows", "success_rows", "failed_rows", "status", "message"])
    return job


def read_rows(path):
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle))
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError("Excel import requires openpyxl. Upload CSV instead.") from exc
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value else "" for cell in next(sheet.iter_rows(max_row=1))]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append({headers[index]: value for index, value in enumerate(row)})
    return rows


def export_dataset(target, file_format="csv"):
    headers, rows = dataset_rows(target)
    if file_format == "xlsx":
        return export_basic_xlsx(f"{target}_export", headers, rows)
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    response = HttpResponse(output.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{target}_export.csv"'
    return response


def dataset_rows(target):
    if target == "books":
        headers = ["Book Code", "Title", "ISBN", "Authors", "Publisher", "Category", "Status"]
        rows = [[book.book_code, book.title, book.isbn or "", book.author_names, book.publisher or "", book.category or "", book.status] for book in Book.objects.prefetch_related("authors").select_related("publisher", "category")]
    elif target == "members":
        headers = ["Member Code", "Name", "Type", "Email", "Status"]
        rows = [[member.member_code, member.display_name, member.get_member_type_display(), member.user.email, member.status] for member in Member.objects.select_related("user")]
    elif target == "borrow_records":
        headers = ["Member", "Book Copy", "Borrow Date", "Due Date", "Return Date", "Status"]
        rows = [[record.member.member_code, record.book_copy.copy_code, record.borrow_date, record.due_date, record.return_date or "", record.status] for record in BorrowRecord.objects.select_related("member", "book_copy")]
    elif target == "fines":
        headers = ["Member", "Amount", "Paid", "Balance", "Reason", "Status"]
        rows = [[fine.member.member_code, fine.amount, fine.paid_amount, fine.balance, fine.reason, fine.status] for fine in Fine.objects.select_related("member")]
    else:
        raise ValueError("Unknown export target.")
    return headers, rows
